#!/usr/bin/env python3
"""
THESIS
The project is a sequence of evidence-driven pivots: physical tracking exposed
latency, VLA experiments exposed action semantics and inconsistency, and a
bounded local IK servo created the final real-time hypothesis.

OWN-WORLD
An engineered editorial system: off-white paper, black ink, measured rules,
Avenir Next for reading, DIN Condensed for numbers/status, and only real-life
robot-camera imagery. No futuristic HUD, no decorative cards, no fake imagery.

STORY
Eleven main slides move from the moving-target problem through system design,
the sim-to-real VLA program, the decision not to deploy, the IK timing shift,
the live demonstration, and the remaining physical question.

FIRST VIEWPORT
A formal research title occupies the left field. A monochrome real XR
demonstration frame fills the right field and is labeled honestly.

FORM
Headline-led 16:9 presentation with one visual thesis per slide, five reactive
appendix slides, scripted four-person handoffs, and a 90-second live demo after
the IK climax. The user-selected direction is deterministic; no seed was used.
"""

from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets" / "real"
OUTPUT = ROOT / "humanoid_robot_grasping_15min.pptx"

SW, SH = 13.333, 7.5

PAPER = "F4F3EF"
WHITE = "FFFFFF"
INK = "0B0B0B"
MID = "64645F"
LIGHT = "D7D6CF"
PALE = "EAE9E3"

AVENIR = "Avenir Next"
DIN = "DIN Condensed"

ROLE = {
    "hero": (50, AVENIR, True, 0.96),
    "title": (36, AVENIR, True, 1.0),
    "takeaway": (30, AVENIR, True, 1.04),
    "body": (21, AVENIR, False, 1.12),
    "body_bold": (21, AVENIR, True, 1.1),
    "diagram": (17, AVENIR, False, 1.08),
    "diagram_bold": (17, AVENIR, True, 1.05),
    "stat": (56, DIN, True, 0.9),
    "stat_small": (38, DIN, True, 0.94),
    "label": (14, DIN, True, 1.0),
    "meta": (11.5, AVENIR, True, 1.0),
    "appendix": (16, AVENIR, False, 1.08),
    "appendix_bold": (16, AVENIR, True, 1.05),
}


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    role: str = "body",
    color: str = INK,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    size: float | None = None,
    font: str | None = None,
    bold: bool | None = None,
    line_spacing: float | None = None,
    margin: float = 0,
    wrap: bool = True,
):
    default_size, default_font, default_bold, default_leading = ROLE[role]
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = valign
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing if line_spacing is not None else default_leading
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = text
    run.font.name = font or default_font
    run.font.size = Pt(size or default_size)
    run.font.bold = default_bold if bold is None else bold
    run.font.color.rgb = rgb(color)
    return box


def add_rich_text(
    slide,
    runs: Sequence[tuple[str, str, str]],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = valign
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = 1.0
    p.space_after = Pt(0)
    for content, role, color in runs:
        size, font, bold, _ = ROLE[role]
        run = p.add_run()
        run.text = content
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return box


def rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: str | None = None,
    line_color: str | None = INK,
    line_width: float = 1,
    dash: bool = False,
):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line_color)
        shape.line.width = Pt(line_width)
        if dash:
            shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return shape


def circle(
    slide,
    x: float,
    y: float,
    d: float,
    *,
    fill: str | None = None,
    line_color: str | None = INK,
    line_width: float = 1,
):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d)
    )
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line_color)
        shape.line.width = Pt(line_width)
    return shape


def line(
    slide,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    *,
    color: str = INK,
    width: float = 1,
    dash: bool = False,
):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    if dash:
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return shape


def arrow(
    slide,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    *,
    color: str = INK,
    width: float = 1.2,
    dash: bool = False,
):
    line(slide, x1, y1, x2, y2, color=color, width=width, dash=dash)
    tri = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE,
        Inches(x2 - 0.06),
        Inches(y2 - 0.06),
        Inches(0.12),
        Inches(0.12),
    )
    if abs(x2 - x1) >= abs(y2 - y1):
        tri.rotation = 90 if x2 >= x1 else 270
    else:
        tri.rotation = 180 if y2 >= y1 else 0
    tri.fill.solid()
    tri.fill.fore_color.rgb = rgb(color)
    tri.line.fill.background()
    return tri


def add_image_crop(slide, path: Path, x: float, y: float, w: float, h: float):
    with Image.open(path) as image:
        iw, ih = image.size
    image_ratio = iw / ih
    frame_ratio = w / h
    picture = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    if image_ratio > frame_ratio:
        visible = frame_ratio / image_ratio
        crop = (1 - visible) / 2
        picture.crop_left = crop
        picture.crop_right = crop
    else:
        visible = image_ratio / frame_ratio
        crop = (1 - visible) / 2
        picture.crop_top = crop
        picture.crop_bottom = crop
    return picture


def tag(slide, text: str, x: float, y: float, w: float, *, dark: bool = False, dashed: bool = False):
    fill = INK if dark else WHITE
    line_color = INK if dark else MID
    rect(slide, x, y, w, 0.36, fill=fill, line_color=line_color, line_width=1, dash=dashed)
    add_text(
        slide,
        text.upper(),
        x + 0.08,
        y + 0.055,
        w - 0.16,
        0.22,
        role="label",
        size=12.5,
        color=WHITE if dark else MID,
        align=PP_ALIGN.CENTER,
    )


def make_slide(prs: Presentation, *, dark: bool = False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(INK if dark else PAPER)
    return slide


def title(slide, text: str, number: int, *, dark: bool = False, size: float | None = None):
    color = WHITE if dark else INK
    add_text(slide, text, 0.7, 0.45, 11.55, 0.7, role="title", size=size, color=color)
    line(slide, 0.7, 1.19, 12.62, 1.19, color=MID if dark else LIGHT, width=0.8)
    add_text(
        slide,
        f"{number:02d}",
        12.2,
        0.55,
        0.4,
        0.24,
        role="meta",
        color=LIGHT if dark else MID,
        align=PP_ALIGN.RIGHT,
    )


def footer(slide, number: int, qualifier: str = "", *, dark: bool = False):
    color = LIGHT if dark else MID
    if qualifier:
        add_text(slide, qualifier, 0.7, 7.08, 10.1, 0.2, role="meta", size=10.8, color=color)
    add_text(
        slide,
        str(number),
        12.2,
        7.05,
        0.4,
        0.2,
        role="meta",
        size=10.8,
        color=color,
        align=PP_ALIGN.RIGHT,
    )


def add_notes(slide, text: str):
    slide.notes_slide.notes_text_frame.text = text.strip()


def build() -> None:
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    prs.core_properties.title = "Real-Time Moving-Object Interception with a Unitree G1"
    prs.core_properties.subject = "UC Santa Cruz Summer AI Academy final presentation"
    prs.core_properties.author = "[Replace with four presenter names]"
    prs.core_properties.comments = (
        "Evidence-backed presentation generated from docs/ and docs/presentation/. "
        "Real images are labeled as XR demonstration data. Replace presenter identity fields."
    )

    frame_060 = ASSETS / "real_xr_demo_060_mono.jpg"
    frame_105 = ASSETS / "real_xr_demo_105_mono.jpg"
    frame_150 = ASSETS / "real_xr_demo_150_mono.jpg"

    # 1 — Title
    s = make_slide(prs)
    add_image_crop(s, frame_105, 7.85, 0, 5.483, 7.5)
    rect(s, 7.85, 0, 0.07, 7.5, fill=INK, line_color=None)
    tag(s, "Tracking demonstrated", 0.72, 0.58, 1.96, dark=True)
    tag(s, "Interception attempt pending", 2.86, 0.58, 2.46, dashed=True)
    add_text(
        s,
        "Real-Time\nMoving-Object\nInterception",
        0.72,
        1.36,
        6.55,
        2.45,
        role="hero",
        size=48,
    )
    add_text(s, "with a Unitree G1", 0.74, 3.87, 5.9, 0.52, role="takeaway", size=29, color=MID)
    add_text(
        s,
        "From VLA experimentation to\nsafety-bounded geometric control",
        0.74,
        4.67,
        6.2,
        0.82,
        role="body",
        size=20,
    )
    add_text(
        s,
        "[Presenter 1 · School]  ·  [Presenter 2 · School]\n"
        "[Presenter 3 · School]  ·  [Presenter 4 · School]",
        0.74,
        5.83,
        6.3,
        0.7,
        role="diagram",
        size=16,
    )
    add_text(s, "UC Santa Cruz Summer AI Academy", 0.74, 6.66, 6.2, 0.28, role="meta", color=MID)
    add_text(
        s,
        "REAL XR DEMONSTRATION DATA",
        8.12,
        6.95,
        4.82,
        0.25,
        role="label",
        size=12.5,
        color=WHITE,
        align=PP_ALIGN.RIGHT,
    )
    footer(s, 1)
    add_notes(
        s,
        """
TIME 0:45 · PRESENTER 1

“This project investigates real-time moving-object interception with a Unitree G1 humanoid. The practical target is controlled palm contact with a moving plush bunny—not a verified fingered grasp. The system combines live perception, prediction, inverse kinematics, trajectory shaping, and robot-side safety. Physical visual tracking has been demonstrated. The optimized interception path is the final bounded hypothesis, and today’s live segment will show guarded tracking followed by one interception attempt.”

The image is a real XR demonstration-data frame, not autonomous execution.

SOURCE: docs/presentation/README.md; CURRENT_STATUS.md
""",
    )

    # 2 — Problem
    s = make_slide(prs)
    title(s, "Moving targets turn latency into the central problem", 2, size=35)
    add_text(s, "THE TARGET MOVES THROUGHOUT THE LOOP", 0.78, 1.53, 4.25, 0.25, role="label", color=MID)
    line(s, 0.82, 2.17, 12.0, 2.17, color=INK, width=1.4)
    target_xs = [1.0, 3.1, 5.25, 7.42, 9.55, 11.7]
    for i, px in enumerate(target_xs):
        circle(s, px, 2.0, 0.34, fill=INK if i == len(target_xs) - 1 else PAPER, line_color=INK, line_width=1.2)
    arrow(s, 1.38, 2.17, 11.52, 2.17, width=1.2)
    stages = [
        ("SEE", "RGB + depth"),
        ("LOCATE", "3D position"),
        ("PREDICT", "crossing"),
        ("MOVE", "arm target"),
        ("VERIFY", "fresh + safe"),
    ]
    for i, (head, body) in enumerate(stages):
        x = 0.82 + i * 2.38
        add_text(s, head, x, 3.02, 1.65, 0.3, role="label")
        add_text(s, body, x, 3.45, 1.8, 0.34, role="diagram", size=16, color=MID)
        if i < len(stages) - 1:
            arrow(s, x + 1.63, 3.27, x + 2.13, 3.27, color=MID)
    line(s, 0.82, 4.26, 12.0, 4.26, color=LIGHT, width=0.9)
    add_text(s, "Controlled palm contact", 0.82, 4.72, 5.1, 0.58, role="takeaway")
    add_text(
        s,
        "must remain inside joint, collision,\ntabletop, freshness, and operator limits.",
        6.1,
        4.74,
        5.6,
        0.92,
        role="body",
        size=20,
        color=MID,
    )
    add_rich_text(
        s,
        [("~57", "stat_small", INK), (" camera FPS     ", "diagram", MID), ("49–51", "stat_small", INK), (" YOLO FPS", "diagram", MID)],
        0.82,
        5.92,
        7.15,
        0.62,
    )
    add_text(s, "representative dry-run perception", 8.45, 6.12, 3.4, 0.26, role="meta", color=MID, align=PP_ALIGN.RIGHT)
    footer(s, 2, "Representative July 27 dry runs; perception speed alone does not establish interception")
    add_notes(
        s,
        """
TIME 1:00 · PRESENTER 1

“Detection alone does not solve interception. The bunny continues moving while a frame is captured, transmitted, detected, fused with depth, converted into a three-dimensional target, and translated into arm motion. That makes end-to-end latency the central problem. The loop must see, locate, predict, move, and continuously verify. Representative dry runs reached about 57 camera frames per second and 49 to 51 YOLO inferences per second, so perception became fast enough to expose the next bottleneck: producing current, safe arm motion before the target moved away.”

SOURCE: docs/presentation/TECHNICAL_EXPLAINER.md; EXPERIMENTS_AND_RESULTS.md
""",
    )

    # 3 — Journey
    s = make_slide(prs)
    title(s, "Each result exposed the next bottleneck", 3)
    phases = [
        ("VISION", "60 FPS-class\ncamera + YOLO", "see reliably"),
        ("PHYSICAL IK", "guarded tracking\ndemonstrated", "move safely"),
        ("SIM + DATA", "9,312 controller trials\n547 episodes", "scale experiments"),
        ("VLA", "36 attempts\n0 authorized", "test learning"),
        ("REAL-TIME PIVOT", "local analytic IK\n+ Ruckig", "attack latency"),
    ]
    y = 2.03
    line(s, 1.1, y, 12.05, y, color=INK, width=1.6)
    for i, (phase, proof, lesson) in enumerate(phases):
        x = 0.82 + i * 2.42
        circle(s, x + 0.3, y - 0.17, 0.34, fill=INK if i == 4 else PAPER, line_color=INK, line_width=1.4)
        add_text(s, phase, x, 2.55, 2.05, 0.3, role="label", color=INK if i == 4 else MID)
        add_text(s, proof, x, 3.0, 2.08, 0.72, role="diagram_bold", size=16.5)
        add_text(s, lesson, x, 3.92, 2.05, 0.28, role="meta", color=MID)
    rect(s, 0.82, 4.78, 11.25, 1.24, fill=INK, line_color=INK)
    add_text(s, "The project did not abandon learning.", 1.13, 5.08, 4.5, 0.42, role="takeaway", size=27, color=WHITE)
    add_text(
        s,
        "It used negative results to choose\nthe shortest defensible live path.",
        6.26,
        5.02,
        5.3,
        0.72,
        role="body",
        size=20,
        color=WHITE,
    )
    footer(s, 3, "Simulation evidence is not hardware validation")
    add_notes(
        s,
        """
TIME 1:05 · PRESENTER 1

“The project evolved through evidence rather than one fixed implementation. First, the team established a stable camera and detector. Next, guarded inverse kinematics moved the physical arm and tracked the bunny, but the response lagged. Simulation then increased experiment throughput, including 9,312 controller trials. A combined real-and-sim dataset made the VLA question testable. Thirty-six VLA attempts exposed problems in target semantics, consistency, and promotion safety. Those findings did not make the learning work irrelevant; they justified a narrower deterministic path based on fast local IK, smooth Ruckig execution, and the safety boundary already demonstrated on hardware.”

[Handoff to Presenter 2]

SOURCE: docs/presentation/PROJECT_JOURNEY.md
""",
    )

    # 4 — Architecture
    s = make_slide(prs)
    title(s, "The robot retains final authority", 4)
    add_text(s, "GB10 · PERCEPTION AND GEOMETRY", 0.82, 1.48, 4.1, 0.26, role="label", color=MID)
    pipeline = [
        ("D435I", "RGB + depth"),
        ("YOLO", "bunny box"),
        ("RGB–D", "3D point"),
        ("α–β", "position + velocity"),
        ("CROSSING", "place + time"),
        ("PINOCCHIO", "joint update"),
        ("CHECKS", "joint · collision · table"),
        ("RUCKIG", "smooth motion"),
    ]
    for i, (head, body) in enumerate(pipeline):
        row = 0 if i < 4 else 1
        col = i if row == 0 else i - 4
        x = 0.82 + col * 2.25
        y = 1.95 + row * 1.65
        rect(s, x, y, 1.78, 0.88, fill=WHITE, line_color=INK, line_width=1.1)
        add_text(s, head, x + 0.12, y + 0.12, 1.54, 0.24, role="label", align=PP_ALIGN.CENTER)
        add_text(s, body, x + 0.1, y + 0.48, 1.58, 0.23, role="meta", size=10.8, color=MID, align=PP_ALIGN.CENTER)
        if col < 3:
            arrow(s, x + 1.83, y + 0.44, x + 2.12, y + 0.44, color=MID)
    arrow(s, 8.82, 2.39, 9.45, 3.95, color=INK, width=1.4)
    rect(s, 9.67, 1.78, 2.55, 3.55, fill=INK, line_color=INK)
    add_text(s, "UNITREE G1", 9.95, 2.12, 1.98, 0.28, role="label", color=LIGHT, align=PP_ALIGN.CENTER)
    add_text(
        s,
        "250 Hz",
        9.82,
        2.76,
        2.23,
        0.72,
        role="stat",
        size=46,
        color=WHITE,
        align=PP_ALIGN.CENTER,
        wrap=False,
    )
    add_text(s, "robot-local\nsafety controller", 9.95, 3.58, 1.98, 0.65, role="diagram_bold", color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, "Only this process publishes\nUnitree arm commands.", 9.92, 4.52, 2.07, 0.55, role="meta", size=11, color=LIGHT, align=PP_ALIGN.CENTER)
    rect(s, 0.82, 5.62, 11.4, 0.65, fill=None, line_color=MID, line_width=1, dash=True)
    add_text(s, "OFFLINE RESEARCH", 1.08, 5.82, 1.72, 0.22, role="label", size=12.5, color=MID)
    add_text(
        s,
        "Dual RTX 3090 · Isaac Sim · MuJoCo/MJX · RLDS/TFDS · UniFoLM-VLA evaluation",
        3.0,
        5.8,
        8.8,
        0.26,
        role="diagram",
        size=15.5,
        color=MID,
    )
    footer(s, 4, "No learned or Cartesian proposal publishes directly to the motors")
    add_notes(
        s,
        """
TIME 1:05 · PRESENTER 2

“The live path separates perception from final hardware authority. On the GB10, the RealSense supplies color and depth, YOLO finds the bunny, RGB-depth fusion creates a three-dimensional point, and an alpha-beta filter estimates position and velocity. A plane-crossing predictor chooses an interception point and time. Pinocchio proposes a joint update; joint, collision, and tabletop checks validate it; Ruckig limits velocity, acceleration, and jerk. The Unitree G1 still owns the final command boundary through a robot-local 250-hertz safety controller. Isaac Sim, MuJoCo, dataset conversion, and VLA evaluation remain offline research infrastructure.”

SOURCE: docs/presentation/TECHNICAL_EXPLAINER.md
""",
    )

    # 5 — Dataset
    s = make_slide(prs)
    title(s, "A sim-to-real dataset made the VLA question testable", 5, size=34)
    photo_paths = [frame_060, frame_105, frame_150]
    for i, path in enumerate(photo_paths):
        x = 0.82 + i * 3.14
        add_image_crop(s, path, x, 1.55, 2.94, 2.05)
    add_text(s, "REAL XR DEMONSTRATION DATA · APPROACH SEQUENCE", 0.82, 3.74, 9.3, 0.24, role="label", size=12.5, color=MID)
    add_rich_text(
        s,
        [
            ("480", "stat", INK),
            (" simulation  +  ", "body", MID),
            ("67", "stat", INK),
            (" real  =  ", "body", MID),
            ("547", "stat", INK),
            (" episodes", "body", MID),
        ],
        0.82,
        4.28,
        10.95,
        0.72,
    )
    line(s, 0.82, 5.18, 12.0, 5.18, color=LIGHT, width=0.9)
    add_text(s, "127", 0.82, 5.52, 0.8, 0.48, role="stat_small", wrap=False)
    add_text(s, "audited", 1.55, 5.66, 0.78, 0.24, role="meta", color=MID)
    arrow(s, 2.55, 5.78, 3.04, 5.78, color=MID)
    add_text(s, "97", 3.23, 5.52, 0.65, 0.48, role="stat_small", wrap=False)
    add_text(s, "accepted", 3.88, 5.66, 0.88, 0.24, role="meta", color=MID)
    arrow(s, 4.97, 5.78, 5.46, 5.78, color=MID)
    add_text(s, "67", 5.66, 5.52, 0.65, 0.48, role="stat_small", wrap=False)
    add_text(s, "canonical real", 6.3, 5.66, 1.38, 0.24, role="meta", color=MID)
    add_text(s, "32,051", 9.16, 5.42, 1.68, 0.52, role="stat_small", align=PP_ALIGN.RIGHT)
    add_text(s, "frames", 10.98, 5.66, 0.82, 0.24, role="meta", color=MID)
    add_text(s, "session-held-out train · validation · sealed test", 8.5, 6.2, 3.3, 0.25, role="meta", size=10.8, color=MID, align=PP_ALIGN.RIGHT)
    footer(s, 5, "Artifact-verified manifests; real frames shown are demonstration data, not autonomous execution")
    add_notes(
        s,
        """
TIME 1:10 · PRESENTER 2

“The VLA program required a common data contract across simulation and real demonstrations. The latest canonical dataset contains 547 episodes and 32,051 frames: 480 generated in Isaac Sim and 67 captured through real XR teleoperation. Real data was deliberately filtered. Of 127 audited episodes, 97 were accepted for materialization and 67 passed the timing, motion, image, and contact contract. Train, validation, and sealed-test splits were separated by capture session to limit leakage risk. The frames shown here are real demonstration data; they illustrate the task and data collection, not autonomous success.”

[Handoff to Presenter 3]

SOURCE: docs/presentation/EXPERIMENTS_AND_RESULTS.md
""",
    )

    # 6 — VLA semantics
    s = make_slide(prs)
    title(s, "Action semantics mattered more than training longer", 6, size=35)
    add_rich_text(
        s,
        [
            ("36", "stat_small", INK),
            (" attempts   →   ", "diagram", MID),
            ("32", "stat_small", INK),
            (" checkpoint-producing   →   ", "diagram", MID),
            ("0", "stat_small", INK),
            (" authorized", "diagram", MID),
        ],
        0.82,
        1.47,
        10.8,
        0.55,
    )
    add_text(s, "WEIGHTED RIGHT-HAND ADE · LOWER IS BETTER", 0.82, 2.26, 3.8, 0.25, role="label", color=MID)
    chart_x, chart_y, max_value = 3.15, 2.82, 0.14
    values = [
        ("absolute action", 0.12517, MID),
        ("relative to current pose", 0.08535, LIGHT),
        ("achieved future state", 0.06695, INK),
    ]
    for i, (name, value, color) in enumerate(values):
        y = chart_y + i * 1.03
        add_text(s, name, 0.82, y + 0.02, 2.1, 0.3, role="diagram", size=16, color=MID if i < 2 else INK)
        width = 6.55 * value / max_value
        rect(s, chart_x, y, width, 0.46, fill=color, line_color=None)
        add_text(s, f"{value:.5f} m", chart_x + width + 0.17, y + 0.04, 1.2, 0.3, role="diagram_bold", size=15.5)
    add_text(
        s,
        "31.8%",
        9.64,
        3.08,
        1.64,
        0.52,
        role="stat_small",
        size=34,
        align=PP_ALIGN.RIGHT,
        wrap=False,
    )
    add_text(s, "absolute → relative", 9.1, 3.66, 2.18, 0.25, role="meta", color=MID, align=PP_ALIGN.RIGHT)
    add_text(
        s,
        "21.6%",
        9.64,
        4.48,
        1.64,
        0.52,
        role="stat_small",
        size=34,
        align=PP_ALIGN.RIGHT,
        wrap=False,
    )
    add_text(s, "relative → future-aligned", 8.78, 5.05, 2.5, 0.25, role="meta", color=MID, align=PP_ALIGN.RIGHT)
    line(s, 0.82, 6.08, 11.18, 6.08, color=INK, width=1.3)
    add_text(s, "Longer training did not improve the original pilot.", 0.82, 6.29, 6.6, 0.36, role="body_bold", size=19)
    add_text(s, "The label definition was the breakthrough.", 7.3, 6.29, 3.98, 0.36, role="body", size=19, color=MID, align=PP_ALIGN.RIGHT)
    footer(s, 6, "Offline evaluation only")
    add_notes(
        s,
        """
TIME 1:15 · PRESENTER 3

“The training store records 36 training or smoke attempts, and 32 produced action checkpoints. The largest improvements came from redefining the action, not extending optimization. Predicting absolute pose produced 0.12517 meters weighted right-hand error. Anchoring motion relative to the measured current pose reduced that error by 31.8 percent. The next experiment aligned the target with the state actually achieved after sensing, inference, transport, and actuation delay. Error fell again, from 0.08535 to 0.06695 meters—a further 21.6 percent. Longer training did not beat the original pilot. The critical discovery was that the model had been trained against the wrong moment in the robot’s motion.”

SOURCE: docs/presentation/VLA_EXPERIMENT_HISTORY.md; VLA_TIMING_ALIGNMENT_RESULTS.md
""",
    )

    # 7 — VLA decision
    s = make_slide(prs)
    title(s, "Lower offline error still produced inconsistent motion", 7, size=34)
    add_text(s, "0", 0.84, 1.48, 3.0, 1.28, role="stat", size=94)
    add_text(s, "robot-authorized\ncheckpoints", 0.88, 2.83, 2.85, 0.72, role="takeaway", size=27)
    line(s, 4.2, 1.55, 4.2, 6.18, color=INK, width=1.2)
    reasons = [
        ("BASELINE", "Did not beat holding the current pose or predicting the mean action."),
        ("2.14×", "Predicted displacement exceeded the real-validation target motion."),
        ("8.8%", "Normalized outputs saturated at the clipping limit."),
    ]
    for i, (head, body) in enumerate(reasons):
        y = 1.58 + i * 1.45
        add_text(s, head, 4.65, y, 1.45, 0.48, role="stat_small", size=34)
        add_text(s, body, 6.25, y + 0.02, 5.45, 0.68, role="body", size=19.5)
        if i < len(reasons) - 1:
            line(s, 4.65, y + 1.02, 11.7, y + 1.02, color=LIGHT, width=0.9)
    rect(s, 4.65, 5.73, 7.05, 0.58, fill=INK, line_color=INK)
    add_text(s, "NOT PROMOTED", 4.88, 5.88, 1.9, 0.25, role="label", color=WHITE)
    add_text(s, "Accuracy, calibration, consistency, and safety all mattered.", 7.02, 5.88, 4.4, 0.25, role="meta", color=LIGHT, align=PP_ALIGN.RIGHT)
    footer(s, 7, "No measured production VLA inference-time claim is used")
    add_notes(
        s,
        """
TIME 1:00 · PRESENTER 3

“The better metric was not enough to justify deployment. The selected model still lost to trivial hold-pose and mean-action baselines. On real validation, its predicted displacement was 2.14 times the target motion, and 8.8 percent of normalized outputs saturated. These failures indicate inconsistency and action-calibration problems that average displacement error alone can hide. The campaign therefore ended with zero robot-authorized checkpoints. This was a deliberate promotion decision: a policy needed to beat baselines, respond to visual input, maintain reasonable action scale, pass simulation, and survive the same deterministic safety gateway used by geometric control.”

[Handoff to Presenter 4]

SOURCE: docs/presentation/EXPERIMENTS_AND_RESULTS.md; VLA_EXPERIMENT_HISTORY.md
""",
    )

    # 8 — IK climax
    s = make_slide(prs, dark=True)
    title(s, "Local servoing changed the timing regime", 8, dark=True)
    tag(s, "Offline canonical benchmark", 0.82, 1.5, 2.5, dark=True)
    add_text(
        s,
        "230.3 ms",
        0.82,
        2.04,
        3.35,
        0.92,
        role="stat",
        size=57,
        color=WHITE,
        wrap=False,
    )
    add_text(s, "bounded global solve", 0.86, 3.05, 2.9, 0.3, role="diagram", color=LIGHT)
    arrow(s, 4.18, 2.57, 5.22, 2.57, color=WHITE, width=1.6)
    add_text(
        s,
        "0.080 ms",
        5.55,
        2.04,
        3.25,
        0.92,
        role="stat",
        size=57,
        color=WHITE,
        wrap=False,
    )
    add_text(s, "bounded analytic local step", 5.6, 3.05, 3.35, 0.3, role="diagram", color=LIGHT)
    add_text(
        s,
        "≈2,879×",
        9.12,
        1.92,
        2.7,
        0.72,
        role="stat",
        size=48,
        color=WHITE,
        align=PP_ALIGN.RIGHT,
        wrap=False,
    )
    add_text(s, "per-call latency gap", 9.22, 2.72, 2.65, 0.3, role="diagram", color=LIGHT, align=PP_ALIGN.RIGHT)
    line(s, 0.82, 3.75, 11.55, 3.75, color=MID, width=1)
    add_text(s, "Different contracts", 0.82, 4.12, 2.55, 0.34, role="label", color=LIGHT)
    add_text(
        s,
        "A global solve searches for a complete pose.\nA local step repeatedly corrects from measured state.",
        3.14,
        4.06,
        5.35,
        0.82,
        role="body",
        size=19,
        color=WHITE,
    )
    add_text(
        s,
        "0.136 → 0.080 ms",
        8.72,
        4.08,
        3.18,
        0.42,
        role="stat_small",
        size=27,
        color=WHITE,
        align=PP_ALIGN.RIGHT,
        wrap=False,
    )
    add_text(s, "41% less local Jacobian compute", 8.75, 4.62, 3.15, 0.28, role="meta", color=LIGHT, align=PP_ALIGN.RIGHT)
    rect(s, 0.82, 5.35, 11.1, 0.94, fill=WHITE, line_color=WHITE)
    add_text(s, "LIVE DEMONSTRATION FOLLOWS", 1.12, 5.64, 3.5, 0.28, role="label", color=INK)
    add_text(s, "guarded tracking  +  one bounded interception attempt", 5.03, 5.59, 6.55, 0.35, role="body_bold", size=20, color=INK, align=PP_ALIGN.RIGHT)
    footer(s, 8, "The ≈2,879× figure is not an end-to-end robot speedup", dark=True)
    add_notes(
        s,
        """
TIME 1:15 + 1:30 LIVE DEMO · PRESENTER 4 / OPERATOR

“The deterministic pivot changed the control contract. The existing bounded global solver required a median 230.3 milliseconds per benchmark target. The new analytic local solver performs one bounded correction in 0.080 milliseconds. That is an approximately 2,879-times per-call latency gap, but it is not a claim that the complete robot became 2,879 times faster: one value is a full global solve and the other is a repeated local servo step. Within the local method itself, replacing a finite-difference Jacobian with Pinocchio’s analytic Jacobian reduced median compute from 0.136 to 0.080 milliseconds, or about 41 percent, while preserving the same one-step error reduction.”

DEMO CUE: “The timing shift made a real-time geometric path plausible. The live demonstration will now show target acquisition, guarded tracking, and one bounded interception attempt.”

DEMO: 20 seconds setup, 60 seconds execution, 10 seconds takeaway. Use assigned operator, spotter, and physical e-stop. If setup misses the rehearsed cutoff, use the real-frame sequence and continue.

SOURCE: docs/G1_IK_BACKEND_EVALUATION.md
""",
    )

    # 9 — Safety after demo
    s = make_slide(prs)
    title(s, "The live path remains bounded by deterministic safety", 9, size=34)
    states = [
        ("ACQUIRE", "stable track"),
        ("PREVIEW", "reachable crossing"),
        ("COMMIT", "revalidated target"),
        ("HOLD", "complete / brief loss"),
    ]
    for i, (head, body) in enumerate(states):
        x = 0.82 + i * 2.67
        fill = INK if i == 2 else WHITE
        rect(s, x, 1.72, 2.05, 0.92, fill=fill, line_color=INK, line_width=1.2)
        add_text(s, head, x + 0.16, 1.91, 1.73, 0.24, role="label", color=WHITE if i == 2 else INK, align=PP_ALIGN.CENTER)
        add_text(s, body, x + 0.12, 2.27, 1.81, 0.22, role="meta", color=LIGHT if i == 2 else MID, align=PP_ALIGN.CENTER)
        if i < len(states) - 1:
            arrow(s, x + 2.13, 2.18, x + 2.52, 2.18, color=MID)
    rect(s, 11.53, 1.72, 0.82, 0.92, fill=None, line_color=MID, line_width=1, dash=True)
    add_text(s, "EXPIRE", 11.57, 1.92, 0.74, 0.23, role="label", size=11.5, color=MID, align=PP_ALIGN.CENTER)
    add_text(s, "stale", 11.57, 2.29, 0.74, 0.2, role="meta", size=9.8, color=MID, align=PP_ALIGN.CENTER)
    add_text(
        s,
        "A target is rejected when\nany boundary fails.",
        0.82,
        3.25,
        6.3,
        0.75,
        role="takeaway",
        size=25,
        line_spacing=0.95,
    )
    checks = [
        "fresh RGB + paired depth",
        "measured arm state",
        "joint + discontinuity limits",
        "self-collision + swept path",
        "table-plane clearance",
        "heartbeat + following error",
    ]
    for i, item in enumerate(checks):
        col = i % 2
        row = i // 2
        x = 0.84 + col * 4.18
        y = 4.22 + row * 0.56
        circle(s, x, y + 0.08, 0.16, fill=INK, line_color=None)
        add_text(s, item, x + 0.34, y, 3.62, 0.28, role="diagram", size=16.5)
    rect(s, 9.12, 3.45, 3.1, 2.37, fill=INK, line_color=INK)
    add_text(s, "250 Hz", 9.45, 3.78, 2.45, 0.65, role="stat", size=50, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, "robot-local guard", 9.45, 4.54, 2.45, 0.28, role="label", color=LIGHT, align=PP_ALIGN.CENTER)
    add_text(s, "network or heartbeat loss\ntriggers a bounded release", 9.43, 5.04, 2.49, 0.54, role="meta", size=11, color=LIGHT, align=PP_ALIGN.CENTER)
    footer(s, 9, "Describe only the behavior actually observed during the live attempt")
    add_notes(
        s,
        """
TIME 1:00 · PRESENTER 3

“The live attempt does not remove the deterministic boundary. The controller first acquires a stable track, previews a reachable crossing, commits only after revalidation, and then holds or expires. Motion is rejected when RGB or depth becomes stale, the arm state is missing, a joint step is too large, collision or table clearance fails, or robot-local following and heartbeat checks fail. The 250-hertz robot controller owns the final release behavior. The demonstration should therefore be interpreted as one bounded physical test, not as proof that every moving-object case is solved.”

Describe the actual demo outcome without converting one attempt into a success rate.

SOURCE: docs/presentation/CURRENT_STATUS.md; TECHNICAL_EXPLAINER.md
""",
    )

    # 10 — Limitations
    s = make_slide(prs)
    title(s, "The remaining question is physical, not computational", 10, size=34)
    columns = [
        ("DEMONSTRATED", ["live perception", "guarded arm movement", "physical bunny tracking"], INK, WHITE),
        ("OFFLINE EVIDENCE", ["547-episode data contract", "VLA ablations", "analytic IK benchmark"], WHITE, INK),
        ("PENDING", ["repeatable fast tracking", "fault-response validation", "bounded interception outcome"], None, INK),
    ]
    for i, (head, items, fill, text_color) in enumerate(columns):
        x = 0.82 + i * 4.02
        rect(s, x, 1.58, 3.55, 3.52, fill=fill, line_color=INK if i < 2 else MID, line_width=1.2, dash=i == 2)
        add_text(s, head, x + 0.22, 1.86, 3.08, 0.28, role="label", color=LIGHT if fill == INK else MID)
        for j, item in enumerate(items):
            y = 2.48 + j * 0.72
            line(s, x + 0.22, y - 0.1, x + 3.28, y - 0.1, color=MID if fill == INK else LIGHT, width=0.7)
            add_text(s, item, x + 0.22, y, 3.0, 0.34, role="diagram_bold", size=16.5, color=text_color)
    add_text(s, "No verified fingered grasp.", 0.82, 5.58, 4.02, 0.4, role="body_bold", size=19)
    add_text(s, "No standardized physical success rate.", 4.84, 5.58, 4.03, 0.4, role="body_bold", size=19)
    add_text(s, "No learned policy authorized.", 8.88, 5.58, 3.24, 0.4, role="body_bold", size=19)
    line(s, 0.82, 6.28, 12.08, 6.28, color=INK, width=1.2)
    add_text(s, "The final hypothesis:", 0.82, 6.47, 2.08, 0.3, role="label")
    add_text(s, "faster local control can preserve tracking quality while reducing enough lag for interception.", 3.02, 6.41, 9.06, 0.42, role="diagram_bold", size=17.5)
    footer(s, 10, "Demonstrated, offline, and pending are intentionally separated")
    add_notes(
        s,
        """
TIME 0:55 · PRESENTER 2

“The evidence classes remain separate. Live perception, guarded arm movement, and physical bunny tracking have been demonstrated. The dataset, VLA ablations, and analytic IK timing are offline evidence. Repeatable fast-path tracking, deliberate fault-response tests, and the bounded interception outcome remain physical validation tasks. The current configuration does not establish a fingered grasp, a standardized physical success rate, or an authorized learned policy. The remaining hypothesis is therefore precise: can the faster local controller preserve the earlier tracking quality while reducing enough end-to-end lag for safe interception?”

[Handoff back to Presenter 1]

SOURCE: docs/presentation/CURRENT_STATUS.md; EVIDENCE_INDEX.md
""",
    )

    # 11 — Conclusion
    s = make_slide(prs, dark=True)
    add_text(s, "Evidence—not optimism—determined the final system", 0.72, 0.62, 11.7, 0.76, role="title", size=38, color=WHITE)
    lines_data = [
        ("PHYSICAL TRACKING", "established that the guarded loop was feasible"),
        ("VLA EXPERIMENTS", "corrected assumptions about actions, timing, and consistency"),
        ("LOCAL CONTROL", "created a practical final real-time hypothesis"),
    ]
    for i, (head, body) in enumerate(lines_data):
        y = 1.7 + i * 1.15
        add_text(s, head, 0.82, y, 2.4, 0.28, role="label", color=LIGHT)
        add_text(s, body, 3.33, y - 0.07, 8.25, 0.48, role="takeaway", size=26, color=WHITE)
        line(s, 0.82, y + 0.63, 11.72, y + 0.63, color=MID, width=0.8)
    add_text(s, "Can it follow fast enough—\nwith evidence and safety?", 0.82, 5.3, 8.75, 1.08, role="takeaway", size=34, color=WHITE)
    add_text(s, "Questions", 10.2, 5.84, 1.55, 0.34, role="body_bold", size=21, color=LIGHT, align=PP_ALIGN.RIGHT)
    add_text(s, "UC Santa Cruz Summer AI Academy", 0.82, 6.82, 5.3, 0.24, role="meta", color=LIGHT)
    add_text(s, "[Replace with four presenter names]", 7.5, 6.82, 4.25, 0.24, role="meta", color=LIGHT, align=PP_ALIGN.RIGHT)
    footer(s, 11, "Main presentation ends here · appendix follows", dark=True)
    add_notes(
        s,
        """
TIME 0:40 · PRESENTER 1

“Three conclusions define the project. Physical tracking established that the guarded perception-to-motion loop was feasible. The VLA program corrected assumptions about how actions, time, and consistency should be represented, while strict promotion gates prevented an unreliable policy from reaching the robot. The analytic local controller then created a practical final hypothesis for real-time interception without removing deterministic safety. The project moved from asking whether the robot could follow the target to asking whether it could follow fast enough—with current evidence and bounded risk.”

“Thank you. We are happy to take questions.”

SOURCE: docs/presentation/README.md; CURRENT_STATUS.md
""",
    )

    # 12 — Appendix technology
    s = make_slide(prs)
    title(s, "Appendix · Technology map", 12)
    entries = [
        ("RealSense D435I", "RGB + per-pixel depth"),
        ("YOLO", "real-time bunny detection"),
        ("RGB–depth fusion", "pixel coordinates → torso-frame 3D"),
        ("Alpha–beta filter", "smoothed position + velocity"),
        ("Plane crossing", "future interception point + time"),
        ("Pinocchio", "G1 kinematics + collision geometry"),
        ("Ruckig", "velocity · acceleration · jerk limits"),
        ("ROS 2 / DDS", "state, depth, commands, services"),
        ("Isaac Sim", "synthetic demonstrations + contact scenes"),
        ("MuJoCo / MJX", "lightweight controller search"),
        ("RLDS / TFDS", "sequential robot-data contract"),
        ("UniFoLM-VLA", "image + instruction + state → proposal"),
    ]
    for i, (head, body) in enumerate(entries):
        col = i % 2
        row = i // 2
        x = 0.82 + col * 6.05
        y = 1.5 + row * 0.85
        add_text(s, head, x, y, 2.15, 0.28, role="appendix_bold")
        add_text(s, body, x + 2.25, y, 3.45, 0.42, role="appendix", color=MID)
        line(s, x, y + 0.57, x + 5.66, y + 0.57, color=LIGHT, width=0.7)
    footer(s, 12, "Reactive Q&A appendix")
    add_notes(s, "Use only when a question asks how a named technology fits the system.")

    # 13 — Appendix VLA inventory
    s = make_slide(prs)
    title(s, "Appendix · The 36 VLA attempts", 13)
    groups = [
        ("Initial full training", "1", "complete adaptation path"),
        ("v28 staged schedule", "5", "simulation → mixed → real"),
        ("Dynamic-action pilots", "4", "23 actions vs right arm vs history"),
        ("Pose23 representation", "6", "absolute vs anchored relative"),
        ("Longer relative training", "2", "more optimization did not help"),
        ("Timing-aligned targets", "2", "one- and three-frame future state"),
        ("Loss-weight controls", "4", "translation, rotation, absolute future"),
        ("v29 motion-focused data", "2", "67 real episodes + moving windows"),
        ("v29 automated campaign", "7", "restarts, seeds, warm starts, gates"),
        ("v30 active horizon", "3", "infrastructure/evaluation failure"),
    ]
    for i, (head, count, outcome) in enumerate(groups):
        col = 0 if i < 5 else 1
        row = i if col == 0 else i - 5
        x = 0.82 + col * 6.0
        y = 1.52 + row * 0.98
        add_text(s, count, x, y, 0.58, 0.42, role="stat_small", size=29)
        add_text(s, head, x + 0.72, y + 0.02, 2.55, 0.28, role="appendix_bold")
        add_text(s, outcome, x + 3.22, y + 0.02, 2.42, 0.48, role="appendix", size=14.5, color=MID)
        line(s, x, y + 0.7, x + 5.65, y + 0.7, color=LIGHT, width=0.7)
    footer(s, 13, "32 attempts produced action checkpoints; 0 were authorized for robot execution")
    add_notes(s, "Use to answer questions about what the 36 attempts actually changed.")

    # 14 — Appendix IK table
    s = make_slide(prs)
    title(s, "Appendix · Canonical IK benchmark", 14)
    headers = ["Backend", "Success / acceptance", "Mean position error", "Median latency", "Unsafe"]
    xs = [0.82, 4.12, 6.55, 8.9, 11.25]
    widths = [3.1, 2.24, 2.1, 1.95, 0.9]
    for x, w, head in zip(xs, widths, headers):
        add_text(s, head, x, 1.52, w, 0.3, role="label", size=12.5, color=MID)
    rows = [
        ("Bounded SciPy global IK", "64 / 64", "1.13 mm", "230.3 ms", "0"),
        ("MuJoCo DLS reference", "0 / 64", "9.53 mm", "15.2 ms", "0*"),
        ("Finite-difference local DLS", "64 / 64 accepted", "0.30 mm / step", "0.136 ms", "0"),
        ("Analytic Pinocchio local DLS", "64 / 64 accepted", "0.30 mm / step", "0.080 ms", "0"),
    ]
    for i, row in enumerate(rows):
        y = 2.03 + i * 0.92
        fill = INK if i == 3 else WHITE
        rect(s, 0.76, y - 0.1, 11.63, 0.72, fill=fill, line_color=INK, line_width=0.8)
        for j, value in enumerate(row):
            add_text(
                s,
                value,
                xs[j],
                y + 0.04,
                widths[j],
                0.33,
                role="appendix_bold" if j in (0, 3) else "appendix",
                size=15,
                color=WHITE if i == 3 else INK,
            )
    add_text(
        s,
        "*No unsafe output in the small interior reference set; the reference did not converge.",
        0.82,
        5.82,
        7.7,
        0.3,
        role="meta",
        color=MID,
    )
    add_text(
        s,
        "Global solves and local servo steps have different contracts.",
        0.82,
        6.28,
        7.7,
        0.35,
        role="body_bold",
        size=18,
    )
    add_text(s, "230.3 ÷ 0.080 = 2,878.75", 8.35, 6.25, 3.8, 0.36, role="stat_small", size=27, align=PP_ALIGN.RIGHT)
    footer(s, 14, "Offline canonical benchmark")
    add_notes(s, "Use to answer methodological questions about the approximately 2,879× per-call comparison.")

    # 15 — Appendix scale
    s = make_slide(prs)
    title(s, "Appendix · Research scale and secondary evidence", 15, size=34)
    stats = [
        ("9,312", "controller candidates evaluated in six hours", "simulation only"),
        ("2.38%", "failure rate in the selected simulation campaign", "not a hardware default"),
        ("22", "attempts before the v29 adaptive campaign paused", "no-progress stop"),
        ("77.6%", "active-motion windows in real training data", "v30 audit"),
        ("754 GB", "datasets, models, runs, and worktrees inspected", "research store"),
        ("25,756", "lines removed from the focused demo branch", "183 files"),
    ]
    for i, (value, body, qualifier) in enumerate(stats):
        col = i % 2
        row = i // 2
        x = 0.82 + col * 6.02
        y = 1.52 + row * 1.54
        add_text(s, value, x, y, 1.7, 0.58, role="stat_small", size=35)
        add_text(s, body, x + 1.82, y + 0.04, 3.65, 0.54, role="appendix_bold")
        add_text(s, qualifier, x + 1.82, y + 0.68, 3.65, 0.24, role="meta", color=MID)
        line(s, x, y + 1.16, x + 5.52, y + 1.16, color=LIGHT, width=0.7)
    footer(s, 15, "Secondary context; these values are not one end-to-end score")
    add_notes(s, "Use only when a question asks about research scale, simulation, or repository focus.")

    # 16 — Appendix evidence and questions
    s = make_slide(prs)
    title(s, "Appendix · Evidence and anticipated questions", 16, size=34)
    evidence = [
        ("DEMONSTRATED", "Observed on physical hardware"),
        ("ARTIFACT-VERIFIED", "Recorded in a manifest or report"),
        ("SOURCE-VERIFIED", "Implementation exists in the repository"),
        ("PENDING", "Implementation exists; physical validation remains"),
    ]
    for i, (head, body) in enumerate(evidence):
        y = 1.48 + i * 0.63
        tag(s, head, 0.82, y, 1.95, dark=i == 0, dashed=i == 3)
        add_text(s, body, 3.05, y + 0.03, 3.9, 0.28, role="appendix", size=15.5)
    line(s, 7.05, 1.48, 7.05, 5.98, color=INK, width=1)
    qa = [
        ("Why not deploy the best VLA?", "It failed baselines, magnitude, consistency, and saturation gates."),
        ("Was IK really 2,879× faster?", "It is a per-call global-solve/local-step gap, not end-to-end speedup."),
        ("Why use simulation?", "It increased coverage without treating simulation as physical proof."),
        ("Why interception, not grasping?", "The current target is controlled palm contact."),
        ("What remains?", "Repeatable fast tracking, induced fault tests, and bounded physical interception."),
    ]
    for i, (question, answer) in enumerate(qa):
        y = 1.48 + i * 0.92
        add_text(s, question, 7.38, y, 2.1, 0.32, role="appendix_bold", size=15.5)
        add_text(s, answer, 9.48, y, 2.82, 0.58, role="appendix", size=14.5, color=MID)
        line(s, 7.38, y + 0.68, 12.25, y + 0.68, color=LIGHT, width=0.7)
    footer(s, 16, "Reactive Q&A appendix")
    add_notes(s, "Use the relevant answer first, then expand from the evidence appendix if requested.")

    prs.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    build()
